#!/usr/bin/env python3
"""Write the preliminary project note as DOCX."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT = Path(__file__).resolve().parents[1] / "docs" / "Porto_Douro_turbidity_preliminary.docx"


def set_run_font(run, size=11, bold=False, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    r = run._element
    rpr = r.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "Calibri")


def add_p(doc, text, *, size=11, bold=False, italic=False, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Calibri"
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        for run in p.runs:
            set_run_font(run)


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("PRELIMINARY RESEARCH NOTE")
    set_run_font(r, size=12, bold=True)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(
        "Flood-driven turbidity in the Douro estuary:\n"
        "a Sentinel-2 C2RCC test against Crestuma discharge"
    )
    set_run_font(r, size=16, bold=True)

    add_p(
        doc,
        "Erasmus mini-project · Universidade do Porto · Geospatial Engineering  ·  "
        "Home university: AGH Krakow, Geoinformatics  ·  Status: concept + first data product",
        italic=True,
        size=10,
    )
    add_p(
        doc,
        "Decision recorded in this note: the optical turbidity project is the one to pursue. "
        "The Sentinel-1 water-extent idea is kept only as optional context, not as the thesis.",
        size=11,
    )

    add_h(doc, "1. Which idea is better — and why", 1)
    add_p(
        doc,
        "Two candidates were compared. Both are local to Porto and both can be processed in SNAP. "
        "They are not equally good as a short, assessable research project.",
    )

    add_h(doc, "Idea A — Sentinel-1 water extent vs rain, tide and dam release", 2)
    add_p(
        doc,
        "Map water with radar (VV) through clouds, then ask whether extra water on the banks is "
        "rainfall, high tide, or a Crestuma–Lever release. Scientifically honest, but weak as a "
        "student thesis: the main channel is always water, so total area barely changes; tide and "
        "discharge dominate the estuary; the result is easy to misread as “rain = flood”. "
        "SNAP is used only as a water mask. There are no optical bands and no C2RCC. "
        "It is a good cautionary appendix, not the main story."
    )

    add_h(doc, "Idea B — Sentinel-2 turbidity / TSM vs Crestuma discharge (CHOSEN)", 2)
    add_p(
        doc,
        "After high river discharge the Douro is known to send a brown sediment plume onto the "
        "shelf. Earlier ocean-colour papers (MODIS/MERIS) already show that this plume is visible "
        "when flow exceeds about 500 m³/s. Sentinel cannot see metals, bacteria or sewage. "
        "It can see mud in the water (TSM / turbidity) with the optical bands, using SNAP C2RCC. "
        "That is exactly the tool a SNAP supervisor expects for water quality, it uses real bands, "
        "and it has a yes/no hypothesis against an independent number (Crestuma Q)."
    )

    add_h(doc, "Why B wins", 2)
    add_bullets(
        doc,
        [
            "SNAP is at the centre (C2RCC / C2X), not only a click-and-threshold mask.",
            "The thesis is one sentence: after high Q the estuary and plume are muddier on cloud-free Sentinel-2.",
            "It is local and already in the Porto literature — we do not pretend to invent the plume; we test it at 10 m, event by event.",
            "It does not collide with the AGH engineering thesis (Kozienice, TIR, thermal power).",
            "Water temperature is dropped from the main thesis: Sentinel-2 has no thermal band; Sentinel-3 is too coarse for this narrow estuary.",
        ],
    )

    add_h(doc, "2. Research problem", 1)
    add_p(
        doc,
        "The Douro estuary (Crestuma–Lever dam to the Atlantic at Porto / Vila Nova de Gaia) is "
        "urban, tidal, and regulated. Eight wastewater plants discharge into the lower basin; "
        "sediments hold trace metals. A flood or a large dam release resuspends and exports fine "
        "sediment. People experience this as dirty water and a plume off Foz / Matosinhos. "
        "A satellite cannot replace a laboratory. It can show whether the optically visible "
        "sediment pulse is larger when Crestuma lets more water through."
    )

    add_h(doc, "3. Thesis and hypotheses", 1)
    add_p(
        doc,
        "Thesis. Cloud-free Sentinel-2 scenes processed in SNAP (C2RCC) show higher turbidity / TSM "
        "in the lower Douro and a larger coastal plume after high Crestuma discharge than on low-discharge days.",
        italic=True,
    )
    add_p(
        doc,
        "H1 (estuary). Median TSM in a fixed lower-estuary polygon is higher on Sentinel-2 dates "
        "1–10 days after Crestuma Q > 500 m³/s than on low-Q dates.",
    )
    add_p(
        doc,
        "H2 (plume). The area (km²) where TSM exceeds a stated threshold, from the mouth onto the "
        "inner shelf, increases with Q.",
    )
    add_p(
        doc,
        "Either accept or reject is a valid result. Clouds after storms will reduce the number of "
        "usable dates; that is a limitation, not a failure. TSM is a proxy for the sediment pulse, "
        "not a concentration of pollutants."
    )

    add_h(doc, "4. What is measured — and what is not", 1)
    add_p(doc, "Measured (optical, SNAP C2RCC):", bold=True)
    add_bullets(
        doc,
        [
            "TSM / turbidity — suspended sediment, the brown signal.",
            "Optional: Kd (light attenuation). Chlorophyll-a only as a cautious appendix (unreliable in muddy water).",
        ],
    )
    add_p(doc, "Not measured (do not claim):", bold=True)
    add_bullets(
        doc,
        [
            "Metals (Pb, Cu, Zn, …), bacteria, sewage markers, nutrients as concentrations.",
            "Water temperature as a main result (no S2 thermal band).",
            "Quality during the storm itself (Sentinel-2 needs a clear sky after the event).",
        ],
    )
    add_p(doc, "Context numbers joined to each scene:", bold=True)
    add_bullets(
        doc,
        [
            "Crestuma–Lever discharge Q (SNIRH) — the flood / release switch.",
            "Tidal height at overpass (Instituto Hidrográfico / Leixões) — so high tide is not mistaken for a flood pulse.",
        ],
    )

    add_h(doc, "5. Data and SNAP method", 1)
    add_p(
        doc,
        "Sensor: Sentinel-2 MSI, Level-1C (top-of-atmosphere). C2RCC is designed for L1C. "
        "Bands used: visible, red-edge and NIR (water colour); SWIR (B11) to help separate water from land and cloud. "
        "Tile of interest: T29TNF (Porto / Douro mouth). Target: 8–15 cloud-free scenes mixing high-Q and low-Q dates."
    )
    add_p(doc, "SNAP graph (same settings on every date):", bold=True)
    add_bullets(
        doc,
        [
            "Read Sentinel-2 L1C",
            "Subset to the Douro mouth + inner shelf",
            "C2RCC (or C2X if the water is extremely turbid)",
            "Export TSM (and the valid-water flags)",
            "Write GeoTIFF / keep the DIMAP for the report appendix",
        ],
    )
    add_p(
        doc,
        "Python after SNAP: clip a fixed estuary polygon and a fixed shelf polygon; compute median TSM "
        "and plume area; join Q and tide; boxplot (H1) and scatter (H2). Two maps only — one low-Q day, "
        "one high-Q day. The decision sits in the table, not in a gallery of screenshots."
    )
    add_p(
        doc,
        "First data product in this repository: a BEAM-DIMAP subset over the Douro / Porto area "
        "(see data/README.md and data/processed/). Full L1C SAFE scenes for the C2RCC series will be "
        "added date by date; the DIMAP is the SNAP-native working format."
    )

    add_h(doc, "6. Calculations (so the report is not “before / after photos”)", 1)
    add_p(doc, "For each usable Sentinel-2 date:")
    add_bullets(
        doc,
        [
            "TSM_est — median TSM in the lower-estuary polygon (avoid mixed bank pixels).",
            "A_plume — km² on the inner shelf above a fixed TSM threshold.",
            "Q — Crestuma discharge on that day (or the 1–3 days before).",
            "η — tidal height at the satellite overpass.",
        ],
    )
    add_p(
        doc,
        "Then: high-Q vs low-Q boxplot of TSM_est (H1); scatter of A_plume versus Q (H2); "
        "one alternative TSM threshold as sensitivity. If H2 disappears when the threshold moves, "
        "that is reported, not hidden."
    )

    add_h(doc, "7. Study area (working box)", 1)
    add_p(
        doc,
        "Approximate AOI: 8.85°W–8.45°W, 41.05°N–41.22°N. It covers the lower estuary (Porto–Gaia), "
        "the mouth, São Paio / the sand spit, and a strip of inner shelf where the plume appears. "
        "The dam at Crestuma–Lever (~21 km upstream) is the hydrological upstream boundary even if "
        "it sits near the edge of the optical subset."
    )

    add_h(doc, "8. Deliverables and limits", 1)
    add_bullets(
        doc,
        [
            "This note + later 8–12 page assessed report organised around H1/H2.",
            "GitHub repository: SNAP graph, Python, figures, BEAM-DIMAP subset.",
            "Short demo in SNAP (C2RCC on one scene).",
        ],
    )
    add_p(
        doc,
        "Limits to state in the report: no in-situ TSM match-ups in this Erasmus window, so comparison "
        "is relative between dates, not “official mg/L for the city”; C2RCC is uncertain in a narrow "
        "urban estuary (adjacency, bottom, sunglint); n is small; we are not the first to see the Douro "
        "plume. Innovation at this level is a repeatable 10 m SNAP test against discharge, not a new sensor."
    )

    add_h(doc, "9. Key references to cite later", 1)
    add_p(
        doc,
        "Mendes et al. (2014, 2017) — Douro turbid plume from ocean-colour imagery; plume typically "
        "detected when river flow exceeds ~500 m³/s. Azevedo et al. (2010) — Crestuma discharge "
        "controls estuary hydrodynamics. Iglesias / Weber de Melo et al. — flood levels driven mainly "
        "by river discharge (and tide), not by local city rain alone. SNAP C2RCC / Brockmann et al. — "
        "atmospheric correction and water constituents for optically complex waters."
    )

    add_h(doc, "10. Immediate next steps", 1)
    add_bullets(
        doc,
        [
            "Open the Porto BEAM-DIMAP in SNAP; inspect RGB and water pixels.",
            "Ask the supervisor: C2RCC vs C2X; number of scenes; deadline.",
            "Collect 8–15 L1C dates (high Q and low Q) and SNIRH discharge.",
            "Run the same C2RCC graph; build the table; decide H1/H2.",
        ],
    )

    add_p(
        doc,
        "This document is a preliminary note, not the assessed report.",
        italic=True,
        size=10,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
